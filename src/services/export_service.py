import os
import tempfile
import time
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.enums import TA_CENTER, TA_LEFT

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from docx import Document
from docx.shared import Pt as DocPt, RGBColor as DocRGB, Inches as DocInches
from docx.enum.text import WD_ALIGN_PARAGRAPH


class ExportService:

    def generate_pdf(self, data: dict) -> str:
        tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False, prefix="testa_")
        doc = SimpleDocTemplate(tmp.name, pagesize=A4,
                                rightMargin=2*cm, leftMargin=2*cm,
                                topMargin=2*cm, bottomMargin=2*cm)
        styles = getSampleStyleSheet()
        story = []

        # Title
        title_style = ParagraphStyle("Title", fontSize=24, textColor=colors.HexColor("#0F172A"),
                                     spaceAfter=6, alignment=TA_CENTER, fontName="Helvetica-Bold")
        subtitle_style = ParagraphStyle("Sub", fontSize=12, textColor=colors.HexColor("#64748B"),
                                        spaceAfter=20, alignment=TA_CENTER)
        story.append(Paragraph("Testa — Automated Testing Report", title_style))
        story.append(Paragraph(f"Generated: {time.strftime('%Y-%m-%d %H:%M UTC', time.gmtime())}", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#3B82F6")))
        story.append(Spacer(1, 20))

        # Summary
        summary = data.get("summary", {})
        h2 = ParagraphStyle("H2", fontSize=14, textColor=colors.HexColor("#1E293B"),
                             fontName="Helvetica-Bold", spaceBefore=12, spaceAfter=8)
        story.append(Paragraph("Test Summary", h2))

        summary_data = [
            ["Metric", "Value"],
            ["Total Tests", str(summary.get("total", 0))],
            ["Passed", str(summary.get("passed", 0))],
            ["Failed", str(summary.get("failed", 0))],
            ["Warnings", str(summary.get("warnings", 0))],
            ["Pass Rate", f"{summary.get('pass_rate', 0)}%"],
            ["Files Analyzed", str(data.get("file_count", 0))],
        ]
        t = Table(summary_data, colWidths=[8*cm, 8*cm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3B82F6")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 11),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("PADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(t)
        story.append(Spacer(1, 20))

        # Suites
        for suite in data.get("suites", []):
            story.append(Paragraph(suite.get("suite", ""), h2))
            suite_data = [["Test", "Status", "Category", "Message"]]
            for test in suite.get("tests", [])[:50]:
                status = test.get("status", "")
                color_map = {"passed": "✅", "failed": "❌", "warning": "⚠️", "skipped": "⏭️"}
                suite_data.append([
                    test.get("name", "")[:40],
                    color_map.get(status, status),
                    test.get("category", ""),
                    test.get("message", "")[:60],
                ])
            t2 = Table(suite_data, colWidths=[5*cm, 2*cm, 3*cm, 7*cm])
            t2.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("WORDWRAP", (0, 0), (-1, -1), True),
            ]))
            story.append(t2)
            story.append(Spacer(1, 15))

        doc.build(story)
        return tmp.name

    def generate_pptx(self, data: dict) -> str:
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        DARK = RGBColor(15, 23, 42)
        BLUE = RGBColor(59, 130, 246)
        GREEN = RGBColor(34, 197, 94)
        RED = RGBColor(239, 68, 68)
        YELLOW = RGBColor(234, 179, 8)
        WHITE = RGBColor(255, 255, 255)
        SLATE = RGBColor(100, 116, 139)

        blank = prs.slide_layouts[6]

        def add_rect(slide, l, t, w, h, color):
            shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
            shape.fill.solid()
            shape.fill.fore_color.rgb = color
            shape.line.fill.background()
            return shape

        def add_text(slide, text, l, t, w, h, size, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
            txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = align
            run = p.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color

        # Slide 1: Title
        slide = prs.slides.add_slide(blank)
        add_rect(slide, 0, 0, 13.33, 7.5, DARK)
        add_rect(slide, 0, 0, 0.3, 7.5, BLUE)
        add_text(slide, "TESTA", 0.6, 1.5, 12, 1.2, 72, bold=True, color=WHITE)
        add_text(slide, "Automated Testing Intelligence Report", 0.6, 2.8, 12, 0.6, 20, color=SLATE)
        add_text(slide, f"Generated {time.strftime('%B %d, %Y', time.gmtime())}", 0.6, 3.5, 12, 0.5, 14, color=SLATE)

        # Summary numbers
        summary = data.get("summary", {})
        stats = [
            (str(summary.get("total", 0)), "TOTAL TESTS", BLUE),
            (str(summary.get("passed", 0)), "PASSED", GREEN),
            (str(summary.get("failed", 0)), "FAILED", RED),
            (f"{summary.get('pass_rate', 0)}%", "PASS RATE", YELLOW),
        ]
        for i, (val, label, color) in enumerate(stats):
            x = 0.6 + i * 3.1
            add_rect(slide, x, 5.2, 2.7, 1.8, RGBColor(30, 41, 59))
            add_text(slide, val, x + 0.1, 5.3, 2.5, 0.9, 36, bold=True, color=color, align=PP_ALIGN.CENTER)
            add_text(slide, label, x + 0.1, 6.2, 2.5, 0.5, 10, color=SLATE, align=PP_ALIGN.CENTER)

        # Slide 2: Suite Overview
        slide2 = prs.slides.add_slide(blank)
        add_rect(slide2, 0, 0, 13.33, 7.5, RGBColor(248, 250, 252))
        add_rect(slide2, 0, 0, 13.33, 1.0, DARK)
        add_text(slide2, "Test Suite Results", 0.4, 0.15, 10, 0.7, 24, bold=True, color=WHITE)

        for i, suite in enumerate(data.get("suites", [])[:6]):
            row = i // 2
            col = i % 2
            x = 0.4 + col * 6.4
            y = 1.3 + row * 1.9
            passed = suite.get("passed", 0)
            total = suite.get("total", 1)
            pct = (passed / total * 100) if total > 0 else 0
            bar_color = GREEN if pct >= 70 else (YELLOW if pct >= 40 else RED)
            add_rect(slide2, x, y, 6.0, 1.6, WHITE)
            add_text(slide2, suite.get("suite", ""), x + 0.15, y + 0.1, 5.7, 0.45, 12, bold=True, color=DARK)
            add_text(slide2, f"P:{suite.get('passed',0)} F:{suite.get('failed',0)} W:{suite.get('warnings',0)} / {total} tests",
                     x + 0.15, y + 0.55, 5.7, 0.35, 10, color=SLATE)
            # Mini progress bar
            bar_w = max(0.1, min(5.5, 5.5 * pct / 100))
            add_rect(slide2, x + 0.15, y + 1.0, 5.5, 0.25, RGBColor(226, 232, 240))
            add_rect(slide2, x + 0.15, y + 1.0, bar_w, 0.25, bar_color)
            add_text(slide2, f"{pct:.0f}%", x + 5.4, y + 0.95, 0.5, 0.35, 10, color=DARK)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False, prefix="testa_")
        prs.save(tmp.name)
        return tmp.name

    def generate_docx(self, data: dict) -> str:
        doc = Document()

        # Title
        title = doc.add_heading("Testa — Automated Testing Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f"Generated: {time.strftime('%Y-%m-%d %H:%M UTC', time.gmtime())}")
        doc.add_paragraph()

        # Summary
        doc.add_heading("Executive Summary", 1)
        summary = data.get("summary", {})
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Metric"
        hdr[1].text = "Value"
        for metric, val in [
            ("Total Tests", summary.get("total", 0)),
            ("Passed", summary.get("passed", 0)),
            ("Failed", summary.get("failed", 0)),
            ("Warnings", summary.get("warnings", 0)),
            ("Pass Rate", f"{summary.get('pass_rate', 0)}%"),
            ("Files Analyzed", data.get("file_count", 0)),
        ]:
            row = table.add_row().cells
            row[0].text = str(metric)
            row[1].text = str(val)
        doc.add_paragraph()

        # Suites
        for suite in data.get("suites", []):
            doc.add_heading(suite.get("suite", ""), 2)
            doc.add_paragraph(
                f"Total: {suite.get('total', 0)} | Passed: {suite.get('passed', 0)} | "
                f"Failed: {suite.get('failed', 0)} | Warnings: {suite.get('warnings', 0)}"
            )
            t = doc.add_table(rows=1, cols=3)
            t.style = "Table Grid"
            t.rows[0].cells[0].text = "Test Name"
            t.rows[0].cells[1].text = "Status"
            t.rows[0].cells[2].text = "Details"
            for test in suite.get("tests", [])[:40]:
                row = t.add_row().cells
                row[0].text = test.get("name", "")[:60]
                row[1].text = test.get("status", "").upper()
                row[2].text = test.get("message", "")[:100]
            doc.add_paragraph()

        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False, prefix="testa_")
        doc.save(tmp.name)
        return tmp.name